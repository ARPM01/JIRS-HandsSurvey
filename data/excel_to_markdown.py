"""Export selected sheets from the Excel workbook to Markdown tables.

Sheets exported:
  'Robotic Hands'                → Robotic_hands.md
  'Included articles (analysis)' → Included_articles.md

Run from grasp_survey_scripts/:
    python data/excel_to_markdown.py
"""
import os

import openpyxl

_DIR = os.path.dirname(os.path.abspath(__file__))
_XLSX = os.path.join(_DIR, "included articles.xlsx")

# (sheet name, header row index, output filename)
_EXPORTS = [
    ("Robotic Hands",                 2, "Robotic_hands.md"),
    ("Included articles (analysis)",  1, "Included_articles.md"),
]

_TITLE_MAX = 60

# Columns dropped from the Included articles table (annotation metadata).
_DROP_COLS = {
    "Comments", "Comments2", "Tactile Location", "Tactile Type",
    "Tactile Spatial Resolution", "Haptic in Skill?",
    "Proprio Angle", "Proprio Torque",
}

# Abbreviated headers for mechanism/sensor columns.
_HEADER_ABBREV = {
    "Num. of Fingers":        "Fingers",
    "Num. Actuators":         "Actuators",
    "Palm (used or not)":     "Palm",
    "Reposition / Opposition":"Repos.",
    "Abduction / Adduction":  "Abd.",
    "Flexion/Extension":      "Flex.",
    "Tactile Feedback":       "Tactile",
    "Kinesthetic Feedback":   "Kinesth.",
    "Visual Feedback":        "Visual",
}

# Maps full skill column header → (category number, skill name for legend).
# DoF columns map to None for the legend (no separate entry).
_SKILL_COLS = {
    "No Contact -> No Motion: Rest position":
        ("Cat. 1",      "Rest position"),
    "No Contact -> Motion -> Not Within Hand: Waving":
        ("Cat. 2",      "Waving"),
    "No Contact -> Motion -> Within Hand: Signing":
        ("Cat. 3",      "Signing"),
    "Contact -> Non-Prehensile -> No Motion -> No Motion at Contact: Open handed hold":
        ("Cat. 4",      "Open handed hold"),
    "Contact -> Non-Prehensile -> No Motion -> Motion at Contact: Open handed hold with external force":
        ("Cat. 5",      "Open handed hold with external force"),
    "Contact -> Non-Prehensile -> Motion -> Not Within Hand -> No Motion at Contact: Pushing coin":
        ("Cat. 6",      "Pushing coin"),
    "Contact -> Non-Prehensile -> Motion -> Not Within Hand -> Motion at Contact: Tactile surface exploration":
        ("Cat. 7",      "Tactile surface exploration"),
    "Contact -> Non-Prehensile -> Motion -> Within Hand -> No Motion at Contact: Flipping light switch":
        ("Cat. 8",      "Flipping light switch"),
    "Contact -> Non-Prehensile -> Motion -> Within Hand -> Motion at Contact: Rolling ball on table":
        ("Cat. 9",      "Rolling ball on table"),
    "Contact -> Prehensile -> No Motion -> No Motion at Contact: Holding object still":
        ("Cat. 10",     "Holding object still"),
    "Contact -> Prehensile -> No Motion -> Motion at Contact: Holding object with external force":
        ("Cat. 11",     "Holding object with external force"),
    "Contact -> Prehensile -> Motion -> Not Within Hand -> No Motion at Contact: Turning doorknob":
        ("Cat. 12",     "Turning doorknob"),
    "Contact -> Prehensile -> Motion -> Not Within Hand -> Motion at Contact: Sliding hand along handrail":
        ("Cat. 13",     "Sliding hand along handrail"),
    "Contact -> Prehensile -> Motion -> Within Hand -> No Motion at Contact: Writing":
        ("Cat. 14",     "Writing"),
    "Contact -> Prehensile -> Motion -> Within Hand -> Motion at Contact: Reorienting pen":
        ("Cat. 15",     "Reorienting pen"),
    "Contact -> Prehensile -> Motion -> Within Hand -> No Motion at Contact: DoF":
        ("Cat. 14 DoF", None),
    "Contact -> Prehensile -> Motion -> Within Hand -> Motion at Contact: DoF":
        ("Cat. 15 DoF", None),
}


def _cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def _truncate(title: str) -> str:
    return title if len(title) <= _TITLE_MAX else title[:_TITLE_MAX].rstrip() + "…"


def sheet_to_markdown(ws, header_row: int, drop_empty_cols: bool = False) -> str:
    rows = [
        [_cell(v) for v in row]
        for row in ws.iter_rows(min_row=header_row, values_only=True)
        if any(v is not None for v in row)
    ]
    if not rows:
        return ""

    if drop_empty_cols:
        n_cols = len(rows[0])
        keep = [i for i in range(n_cols) if any(row[i] for row in rows)]
        rows = [[row[i] for i in keep] for row in rows]

    headers = rows[0]
    sep = ["---"] * len(headers)
    lines = (
        ["| " + " | ".join(headers) + " |"]
        + ["| " + " | ".join(sep) + " |"]
        + ["| " + " | ".join(row) + " |" for row in rows[1:]]
    )
    return "\n".join(lines)


def included_articles_to_markdown(ws) -> str:
    rows = [
        list(row)
        for row in ws.iter_rows(min_row=1, values_only=True)
        if any(v is not None for v in row)
    ]
    if not rows:
        return ""

    headers = [_cell(v) for v in rows[0]]
    papers_i = headers.index("Papers")
    link_i   = headers.index("Link")

    # Identify skill and DoF column indices
    skill_indices = {i for i, h in enumerate(headers) if h in _SKILL_COLS}
    # Indices to skip entirely (Link col + dropped annotation cols + skill cols
    # — skills are collapsed into one "Skills" cell)
    skip = {link_i} | {i for i, h in enumerate(headers) if h in _DROP_COLS} | skill_indices

    # Build new headers: Paper first, then mechanism/sensor cols, then Skills
    new_headers = []
    skills_inserted = False
    for i, h in enumerate(headers):
        if i in skip:
            continue
        if i == papers_i:
            new_headers.append("Paper")
        else:
            if not skills_inserted and i > max(skill_indices):
                new_headers.append("Skills")
                skills_inserted = True
            new_headers.append(_HEADER_ABBREV.get(h, h))
    if not skills_inserted:
        new_headers.append("Skills")

    # Build data rows
    new_rows = []
    for raw in rows[1:]:
        if not any(v is not None for v in raw):
            continue

        # Paper cell
        title = _cell(raw[papers_i])
        url   = _cell(raw[link_i]) if raw[link_i] else ""
        paper_cell = f"[{_truncate(title)}]({url})" if url else _truncate(title)

        # Skills cell: list active category numbers + DoF values
        skill_parts = []
        for i, h in enumerate(headers):
            if i not in skill_indices:
                continue
            v = raw[i]
            cat, legend_name = _SKILL_COLS[h]
            if legend_name is None:          # DoF column
                s = _cell(v)
                if s and s not in ("-", ""):
                    skill_parts.append(f"{cat}:{s}")
            else:                            # boolean skill column
                if str(v).strip().lower() in ("yes", "true", "1"):
                    skill_parts.append(cat.replace("Cat. ", ""))
        skills_cell = ", ".join(skill_parts) if skill_parts else "—"

        row_cells = []
        skills_inserted = False
        for i, h in enumerate(headers):
            if i in skip:
                continue
            if i == papers_i:
                row_cells.append(paper_cell)
            else:
                if not skills_inserted and i > max(skill_indices):
                    row_cells.append(skills_cell)
                    skills_inserted = True
                row_cells.append(_cell(raw[i]))
        if not skills_inserted:
            row_cells.append(skills_cell)

        new_rows.append(row_cells)

    sep = ["---"] * len(new_headers)
    table_lines = (
        ["| " + " | ".join(new_headers) + " |"]
        + ["| " + " | ".join(sep) + " |"]
        + ["| " + " | ".join(row) + " |" for row in new_rows]
    )

    # Category legend
    legend_lines = [
        "",
        "## Category legend",
        "",
        "| Cat. | Skill |",
        "| --- | --- |",
    ] + [
        f"| {short.replace('Cat. ', '')} | {name} |"
        for short, name in _SKILL_COLS.values()
        if name is not None
    ]

    return "\n".join(table_lines) + "\n" + "\n".join(legend_lines)


def main():
    wb = openpyxl.load_workbook(_XLSX, data_only=True)

    for sheet_name, header_row, out_name in _EXPORTS:
        ws = wb[sheet_name]
        if sheet_name == "Included articles (analysis)":
            md = included_articles_to_markdown(ws)
        else:
            md = sheet_to_markdown(ws, header_row, drop_empty_cols=True)
        out_path = os.path.join(_DIR, out_name)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(f"# {sheet_name}\n\n{md}\n")
        print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
